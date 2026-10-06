from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse


def exportar_excel(nombre_archivo, cabeceras, filas, titulo_hoja="Datos"):
    """
    Genera un archivo Excel con cabeceras y filas.
    
    - nombre_archivo: nombre del archivo .xlsx (ej: 'Facturas_20261006.xlsx')
    - cabeceras: lista de strings (nombres de columnas)
    - filas: lista de listas (cada fila con los valores)
    - titulo_hoja: nombre de la pestaña (por defecto 'Datos')
    """
    wb = Workbook()
    ws = wb.active
    ws.title = titulo_hoja
    
    # Estilo de cabeceras
    color_cabecera = "2d6a4f"  # Verde corporativo
    fuente_cabecera = Font(bold=True, color="FFFFFF", size=11)
    fill_cabecera = PatternFill(
        start_color=color_cabecera,
        end_color=color_cabecera,
        fill_type="solid"
    )
    alineacion_centro = Alignment(horizontal="center", vertical="center")
    
    # Escribir cabeceras
    for col_num, cabecera in enumerate(cabeceras, 1):
        celda = ws.cell(row=1, column=col_num, value=cabecera)
        celda.font = fuente_cabecera
        celda.fill = fill_cabecera
        celda.alignment = alineacion_centro
    
    # Escribir filas
    for row_num, fila in enumerate(filas, 2):
        for col_num, valor in enumerate(fila, 1):
            ws.cell(row=row_num, column=col_num, value=valor)
    
    # Auto-ancho de columnas
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_length + 4, 10)
    
    # Fijar la fila de cabeceras
    ws.freeze_panes = "A2"
    
    # Respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response