from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.db.models import Sum, Q
from datetime import datetime, timedelta, date
from weasyprint import HTML

from .models import (
    Factura, CuentaPorPagar, PagoCuentaPorPagar,
    CuentaPorCobrar, CobroCuentaPorCobrar
)
from clientes.models import Cliente
from proveedores.models import Proveedor
from pedidos.models import Pedido
from core.models import DatosEmpresa
from django.db import transaction
from decimal import Decimal
from pedidos.models import Comercial, Comision
from django.db.models import Sum



# ============================================
# FUNCIÓN AUXILIAR: Crear cuenta para factura
# ============================================

def _crear_cuenta_para_factura(factura):
    """
    Crea la CuentaPorCobrar (si INGRESO) o CuentaPorPagar (si GASTO)
    asociada a una factura. No hace nada si:
    - Ya existe una cuenta con ese numero_factura
    - No hay cliente/proveedor en la factura
    """
    # Evitar duplicados
    if factura.tipo == 'INGRESO':
        if CuentaPorCobrar.objects.filter(numero_factura=factura.numero_factura).exists():
            return None
        if not factura.cliente:
            return None
        
        return CuentaPorCobrar.objects.create(
            concepto=f"Factura {factura.numero_factura}",
            cliente=factura.cliente,
            importe_total=factura.total,
            fecha_emision=factura.fecha,
            fecha_vencimiento=factura.fecha_vencimiento or factura.fecha,
            numero_factura=factura.numero_factura,
            notas=f"Generada desde factura {factura.numero_factura}",
        )
    
    elif factura.tipo == 'GASTO':
        if CuentaPorPagar.objects.filter(numero_factura=factura.numero_factura).exists():
            return None
        if not factura.proveedor:
            return None
        
        return CuentaPorPagar.objects.create(
            concepto=f"Factura {factura.numero_factura}",
            categoria='PROVEEDOR',
            proveedor=factura.proveedor,
            importe_total=factura.total,
            fecha_emision=factura.fecha,
            fecha_vencimiento=factura.fecha_vencimiento or factura.fecha,
            numero_factura=factura.numero_factura,
            notas=f"Generada desde factura {factura.numero_factura}",
        )
    
    return None

# ============================================
# FACTURAS
# ============================================

@login_required
def lista_facturas(request):
    # Ver anuladas: por defecto NO
    ver_anuladas = request.GET.get('ver_anuladas', '') == '1'
    
    # 🆕 Filtro por Categoría
    categoria = request.GET.get('categoria', '')

    if ver_anuladas:
        facturas = Factura.objects.all().order_by('-fecha')
    else:
        facturas = Factura.objects.exclude(estado_pago='ANULADO').order_by('-fecha')

    # 🆕 Aplicar filtro de categoría
    if categoria:
        facturas = facturas.filter(categoria=categoria)

    totales = facturas.aggregate(
        ingresos=Sum('total', filter=Q(tipo='INGRESO')),
        gastos=Sum('total', filter=Q(tipo='GASTO')),
    )
    total_ingresos = totales['ingresos'] or 0
    total_gastos = totales['gastos'] or 0
    beneficio = total_ingresos - total_gastos

    context = {
        'facturas': facturas,
        'total_ingresos': total_ingresos,
        'total_gastos': total_gastos,
        'beneficio': beneficio,
        'ver_anuladas': ver_anuladas,
        'categorias': Factura.CATEGORIAS, # 🆕 Enviamos las opciones al HTML
        'categoria_filtro': categoria,    # 🆕
    }
    return render(request, 'finanzas/lista.html', context)


@login_required
def nueva_factura(request):
    """Formulario para crear una nueva factura"""
    if request.method == 'POST':
        try:
            tipo_factura = request.POST.get('tipo')
            cliente_id = request.POST.get('cliente') if tipo_factura == 'INGRESO' else None
            proveedor_id = request.POST.get('proveedor') if tipo_factura == 'GASTO' else None

            factura = Factura.objects.create(
                numero_factura=request.POST.get('numero_factura'),
                numero_albaran=request.POST.get('numero_albaran', ''), # 🆕
                tipo=tipo_factura,
                fecha=request.POST.get('fecha') or datetime.now().date(),
                fecha_vencimiento=request.POST.get('fecha_vencimiento') or None,
                cliente_id=cliente_id or None,
                proveedor_id=proveedor_id or None,
                pedido_id=request.POST.get('pedido') or None,
                categoria=request.POST.get('categoria', 'OTROS'),
                base=request.POST.get('base', 0),
                iva_porcentaje=request.POST.get('iva_porcentaje', 21),
                estado_pago=request.POST.get('estado_pago', 'PENDIENTE'),
                notas=request.POST.get('notas', ''),
                archivo_justificante=request.FILES.get('archivo_factura') or None,
            )
            # Crear cuenta asociada si se ha marcado el checkbox
            if request.POST.get('crear_cuenta'):
                _crear_cuenta_para_factura(factura)

            messages.success(request, f'✅ Factura {factura.numero_factura} creada correctamente')
            return redirect('lista_facturas')

        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    context = {
        'clientes': Cliente.objects.all().order_by('nombre'),
        'proveedores': Proveedor.objects.all().order_by('nombre'),
        'proveedores_json': list(Proveedor.objects.values('id', 'nombre', 'tipo')), # 🆕 Para el JS
        'pedidos': Pedido.objects.all().order_by('-fecha')[:50],
        'tipos': Factura.TIPO_FACTURA,
        'categorias': Factura.CATEGORIAS,
        'estados': Factura.ESTADO_PAGO,
        'fecha_actual': datetime.now().strftime('%Y-%m-%d'),
    }
    return render(request, 'finanzas/nueva_factura.html', context)


@login_required
def editar_factura(request, factura_id):
    """Formulario para editar una factura existente"""
    factura = get_object_or_404(Factura, id=factura_id)
    url_retorno = request.GET.get('return_url', '') # 🆕 Capturamos

    if request.method == 'POST':
        try:
            tipo_factura = request.POST.get('tipo')

            numero_antiguo = factura.numero_factura
            factura.numero_factura = request.POST.get('numero_factura')
            factura.numero_albaran = request.POST.get('numero_albaran', '') # 🆕
            factura.tipo = tipo_factura
            factura.fecha = request.POST.get('fecha')
            factura.fecha_vencimiento = request.POST.get('fecha_vencimiento') or None
            factura.cliente_id = request.POST.get('cliente') if tipo_factura == 'INGRESO' else None
            factura.proveedor_id = request.POST.get('proveedor') if tipo_factura == 'GASTO' else None
            factura.pedido_id = request.POST.get('pedido') or None
            factura.categoria = request.POST.get('categoria')
            factura.base = request.POST.get('base', 0)
            factura.iva_porcentaje = request.POST.get('iva_porcentaje', 21)
            factura.estado_pago = request.POST.get('estado_pago')
            factura.notas = request.POST.get('notas', '')
            # Actualizar archivo si se ha subido uno nuevo
            if request.FILES.get('archivo_factura'):
                factura.archivo_factura = request.FILES.get('archivo_factura')

            factura.save()
            # Sincronizar el nuevo número de factura en las cuentas asociadas (incluido el concepto)
            if numero_antiguo != factura.numero_factura:
                CuentaPorPagar.objects.filter(numero_factura=numero_antiguo).update(
                    numero_factura=factura.numero_factura,
                    concepto=f"Factura {factura.numero_factura}"
                )
                CuentaPorCobrar.objects.filter(numero_factura=numero_antiguo).update(
                    numero_factura=factura.numero_factura,
                    concepto=f"Factura {factura.numero_factura}"
    )


            messages.success(request, f'✅ Factura {factura.numero_factura} actualizada')
            # 🆕 Redirigir con los filtros si existen
            if url_retorno:
                return redirect(f"/finanzas/?{url_retorno}")
            return redirect('lista_facturas')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')


    context = {
        'factura': factura,
        'clientes': Cliente.objects.all().order_by('nombre'),
        'proveedores': Proveedor.objects.all().order_by('nombre'),
        'proveedores_json': list(Proveedor.objects.values('id', 'nombre', 'tipo')), # 🆕 Para el JS
        'pedidos': Pedido.objects.all().order_by('-fecha')[:50],
        'tipos': Factura.TIPO_FACTURA,
        'categorias': Factura.CATEGORIAS,
        'estados': Factura.ESTADO_PAGO,
    }
    return render(request, 'finanzas/editar_factura.html', context)


@login_required
def confirmar_eliminar_factura(request, factura_id):
    factura = get_object_or_404(Factura, id=factura_id)
    context = {
        'factura': factura,
    }
    return render(request, 'finanzas/confirmar_eliminar.html', context)


@login_required
def eliminar_factura(request, factura_id):
    factura = get_object_or_404(Factura, id=factura_id)

    # Solo se pueden eliminar facturas ANULADAS
    if factura.estado_pago != 'ANULADO':
        messages.error(request, '❌ Solo se pueden eliminar facturas ANULADAS. Anúlala primero.')
        return redirect('lista_facturas')

    if request.method == 'POST':
        numero = factura.numero_factura

        with transaction.atomic():
            # Eliminar cuentas asociadas (con sus pagos/cobros por CASCADE)
            CuentaPorCobrar.objects.filter(numero_factura=numero).delete()
            CuentaPorPagar.objects.filter(numero_factura=numero).delete()

            # Eliminar la factura
            factura.delete()

        messages.success(request, f'✅ Factura {numero} y sus cuentas asociadas eliminadas')

        return redirect('lista_facturas')

    return redirect('confirmar_eliminar_factura', factura_id=factura_id)


@login_required
def factura_pdf(request, factura_id):
    """Genera un PDF profesional de la factura"""
    factura = get_object_or_404(Factura, id=factura_id)

    try:
        empresa = DatosEmpresa.objects.first()
    except DatosEmpresa.DoesNotExist:
        empresa = None

    html_string = render_to_string('finanzas/factura_pdf.html', {
        'factura': factura,
        'empresa': empresa,
        'fecha_actual': datetime.now(),
    })

    html = HTML(string=html_string, base_url=request.build_absolute_uri())
    pdf = html.write_pdf()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Factura_{factura.numero_factura}.pdf"'
    return response


# ============================================
# CUENTAS POR PAGAR
# ============================================

@login_required
def cuentas_por_pagar(request):
    """Lista de cuentas por pagar (deudas)"""
    # 🆕 Ordenación
    orden = request.GET.get('orden', '-fecha_vencimiento') # por defecto: más reciente a más antiguo
    categoria = request.GET.get('categoria', '')
    estado = request.GET.get('estado', '')

    cuentas = CuentaPorPagar.objects.exclude(estado='ANULADO').order_by(orden)

    if categoria:
        cuentas = cuentas.filter(categoria=categoria)
    if estado:
        cuentas = cuentas.filter(estado=estado)

    pendientes = cuentas.filter(estado__in=['PENDIENTE', 'PARCIAL'])
    vencidas = [c for c in pendientes if c.esta_vencida]

    total_pendiente = sum(c.saldo_pendiente for c in pendientes)
    total_vencido = sum(c.saldo_pendiente for c in vencidas)

    context = {
        'cuentas': cuentas,
        'pendientes': pendientes,
        'vencidas': vencidas,
        'total_pendiente': total_pendiente,
        'total_vencido': total_vencido,
        'categorias': CuentaPorPagar.CATEGORIA_CHOICES,
        'estados': CuentaPorPagar.ESTADO_CHOICES,
        'categoria_filtro': categoria,
        'estado_filtro': estado,
        'orden_filtro': orden, # 🆕
    }
    return render(request, 'finanzas/cuentas_por_pagar.html', context)

@login_required
def nueva_cuenta_por_pagar(request):
    """Crear una nueva cuenta por pagar"""
    if request.method == 'POST':
        try:
            cuenta = CuentaPorPagar.objects.create(
                concepto=request.POST.get('concepto'),
                categoria=request.POST.get('categoria'),
                proveedor_id=request.POST.get('proveedor') or None,
                importe_total=request.POST.get('importe_total'),
                fecha_emision=request.POST.get('fecha_emision'),
                fecha_vencimiento=request.POST.get('fecha_vencimiento'),
                periodicidad=request.POST.get('periodicidad', 'UNICO'),
                numero_factura=request.POST.get('numero_factura', ''),
                notas=request.POST.get('notas', ''),
                archivo_factura=request.FILES.get('archivo_factura') or None,
            )
            messages.success(request, '✅ Cuenta por pagar creada')
            return redirect('cuentas_por_pagar')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    context = {
        'categorias': CuentaPorPagar.CATEGORIA_CHOICES,
        'periodicidades': CuentaPorPagar.PERIODICIDAD,
        'proveedores': Proveedor.objects.all().order_by('nombre'),
        'fecha_actual': date.today().strftime('%Y-%m-%d'),
    }
    return render(request, 'finanzas/nueva_cuenta_por_pagar.html', context)


@login_required
def registrar_pago_cuenta(request, cuenta_id):
    """Registrar un pago parcial o total"""
    cuenta = get_object_or_404(CuentaPorPagar, id=cuenta_id)

    if request.method == 'POST':
        try:
            PagoCuentaPorPagar.objects.create(
                cuenta=cuenta,
                importe=request.POST.get('importe'),
                fecha=request.POST.get('fecha'),
                metodo=request.POST.get('metodo', ''),
                referencia=request.POST.get('referencia', ''),
                notas=request.POST.get('notas', ''),
                archivo_justificante=request.FILES.get('archivo_justificante') or None,
            )
            messages.success(request, f'✅ Pago registrado: {request.POST.get("importe")}€')
            return redirect('cuentas_por_pagar')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    context = {
        'cuenta': cuenta,
        'fecha_actual': date.today().strftime('%Y-%m-%d'),
    }
    return render(request, 'finanzas/registrar_pago_cuenta.html', context)


# ============================================
# CUENTAS POR COBRAR
# ============================================

@login_required
def cuentas_por_cobrar(request):
    """Lista de cuentas por cobrar"""
    cuentas = CuentaPorCobrar.objects.exclude(estado='ANULADO').order_by('fecha_vencimiento')

    estado = request.GET.get('estado', '')
    if estado:
        cuentas = cuentas.filter(estado=estado)

    pendientes = cuentas.filter(estado__in=['PENDIENTE', 'PARCIAL'])
    vencidas = [c for c in pendientes if c.estado != 'COBRADO' and c.dias_para_vencer < 0]

    total_pendiente = sum(c.saldo_pendiente for c in pendientes)
    total_vencido = sum(c.saldo_pendiente for c in vencidas)

    context = {
        'cuentas': cuentas,
        'pendientes': pendientes,
        'vencidas': vencidas,
        'total_pendiente': total_pendiente,
        'total_vencido': total_vencido,
        'estados': CuentaPorCobrar.ESTADO_CHOICES,
        'estado_filtro': estado,
    }
    return render(request, 'finanzas/cuentas_por_cobrar.html', context)


@login_required
def nueva_cuenta_por_cobrar(request):
    """Crear una nueva cuenta por cobrar"""
    if request.method == 'POST':
        try:
            cuenta = CuentaPorCobrar.objects.create(
                concepto=request.POST.get('concepto'),
                cliente_id=request.POST.get('cliente'),
                importe_total=request.POST.get('importe_total'),
                fecha_emision=request.POST.get('fecha_emision'),
                fecha_vencimiento=request.POST.get('fecha_vencimiento'),
                numero_factura=request.POST.get('numero_factura', ''),
                notas=request.POST.get('notas', ''),
            )
            messages.success(request, '✅ Cuenta por cobrar creada')
            return redirect('cuentas_por_cobrar')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    context = {
        'clientes': Cliente.objects.all().order_by('nombre'),
        'fecha_actual': date.today().strftime('%Y-%m-%d'),
    }
    return render(request, 'finanzas/nueva_cuenta_por_cobrar.html', context)


@login_required
def registrar_cobro_cuenta(request, cuenta_id):
    """Registrar un cobro"""
    cuenta = get_object_or_404(CuentaPorCobrar, id=cuenta_id)

    if request.method == 'POST':
        try:
            CobroCuentaPorCobrar.objects.create(
                cuenta=cuenta,
                importe=request.POST.get('importe'),
                fecha=request.POST.get('fecha'),
                metodo=request.POST.get('metodo', ''),
                referencia=request.POST.get('referencia', ''),
                notas=request.POST.get('notas', ''),
                archivo_justificante=request.FILES.get('archivo_justificante') or None,
            )
            messages.success(request, '✅ Cobro registrado')
            return redirect('cuentas_por_cobrar')
        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    context = {
        'cuenta': cuenta,
        'fecha_actual': date.today().strftime('%Y-%m-%d'),
    }
    return render(request, 'finanzas/registrar_cobro_cuenta.html', context)


# ============================================
# DASHBOARD FINANCIERO
# ============================================

@login_required
def dashboard_financiero(request):
    """Dashboard con resumen de deudas y cobros"""

    pagar_pendientes = CuentaPorPagar.objects.exclude(estado__in=['PAGADO', 'ANULADO'])
    total_pagar = sum(c.saldo_pendiente for c in pagar_pendientes)
    vencidas_pagar = [c for c in pagar_pendientes if c.esta_vencida]
    total_vencido_pagar = sum(c.saldo_pendiente for c in vencidas_pagar)

    cobrar_pendientes = CuentaPorCobrar.objects.exclude(estado__in=['COBRADO', 'ANULADO'])
    total_cobrar = sum(c.saldo_pendiente for c in cobrar_pendientes)
    vencidas_cobrar = [c for c in cobrar_pendientes if c.dias_para_vencer < 0]
    total_vencido_cobrar = sum(c.saldo_pendiente for c in vencidas_cobrar)

    balance = total_cobrar - total_pagar

    categorias = {}
    for c in pagar_pendientes:
        cat = c.get_categoria_display()
        if cat not in categorias:
            categorias[cat] = 0
        categorias[cat] += float(c.saldo_pendiente)

    context = {
        'total_pagar': total_pagar,
        'total_vencido_pagar': total_vencido_pagar,
        'num_pagar_pendientes': pagar_pendientes.count(),
        'num_vencidas_pagar': len(vencidas_pagar),

        'total_cobrar': total_cobrar,
        'total_vencido_cobrar': total_vencido_cobrar,
        'num_cobrar_pendientes': cobrar_pendientes.count(),
        'num_vencidas_cobrar': len(vencidas_cobrar),

        'balance': balance,
        'categorias_pagar': categorias,

        'proximas_pagar': pagar_pendientes.order_by('fecha_vencimiento')[:5],
        'proximas_cobrar': cobrar_pendientes.order_by('fecha_vencimiento')[:5],
    }
    return render(request, 'finanzas/dashboard_financiero.html', context)


# ============================================
# REGISTRAR COBRO/PAGO DESDE FACTURA
# ============================================

@login_required
def registrar_cobro_pago_factura(request, factura_id):
    """Registra un cobro (si es ingreso) o pago (si es gasto) desde una factura"""
    factura = get_object_or_404(Factura, id=factura_id)

    if factura.tipo == 'INGRESO':
        cuenta_existente = CuentaPorCobrar.objects.filter(numero_factura=factura.numero_factura).first()
    else:
        cuenta_existente = CuentaPorPagar.objects.filter(numero_factura=factura.numero_factura).first()

    if cuenta_existente:
        messages.warning(request, '⚠️ Esta factura ya tiene una cuenta asociada')
        if factura.tipo == 'INGRESO':
            return redirect('cuentas_por_cobrar')
        else:
            return redirect('cuentas_por_pagar')

    if request.method == 'POST':
        try:
            fecha_vencimiento = request.POST.get('fecha_vencimiento')

            if factura.tipo == 'INGRESO':
                cuenta = CuentaPorCobrar.objects.create(
                    concepto=f"Factura {factura.numero_factura}",
                    cliente=factura.cliente,
                    importe_total=factura.total,
                    fecha_emision=factura.fecha,
                    fecha_vencimiento=fecha_vencimiento,
                    numero_factura=factura.numero_factura,
                    notas=f"Generada desde la factura {factura.numero_factura}",
                )
                messages.success(request, '✅ Cuenta por cobrar creada')
                return redirect('registrar_cobro_cuenta', cuenta_id=cuenta.id)
            else:
                cuenta = CuentaPorPagar.objects.create(
                    concepto=f"Factura {factura.numero_factura}",
                    categoria='PROVEEDOR',
                    proveedor=factura.proveedor,
                    importe_total=factura.total,
                    fecha_emision=factura.fecha,
                    fecha_vencimiento=fecha_vencimiento,
                    numero_factura=factura.numero_factura,
                    notas=f"Generada desde la factura {factura.numero_factura}",
                )
                messages.success(request, '✅ Cuenta por pagar creada')
                return redirect('registrar_pago_cuenta', cuenta_id=cuenta.id)

        except Exception as e:
            messages.error(request, f'❌ Error: {str(e)}')

    fecha_sugerida = (factura.fecha + timedelta(days=30)).strftime('%Y-%m-%d')

    context = {
        'factura': factura,
        'fecha_sugerida': fecha_sugerida,
    }
    return render(request, 'finanzas/registrar_cobro_pago_factura.html', context)

# ============================================
# ELIMINAR CUENTAS POR PAGAR / COBRAR
# ============================================

@login_required
def confirmar_eliminar_cuenta_por_pagar(request, cuenta_id):
    """Muestra página de confirmación para eliminar una cuenta por pagar"""
    cuenta = get_object_or_404(CuentaPorPagar, id=cuenta_id)

    # Comprobar si tiene factura asociada
    tiene_factura = Factura.objects.filter(numero_factura=cuenta.numero_factura).exists()

    context = {
        'cuenta': cuenta,
        'tiene_factura': tiene_factura,
    }
    return render(request, 'finanzas/confirmar_eliminar_cuenta_pagar.html', context)


@login_required
def eliminar_cuenta_por_pagar(request, cuenta_id):
    """Elimina definitivamente una cuenta por pagar (solo si no tiene factura)"""
    cuenta = get_object_or_404(CuentaPorPagar, id=cuenta_id)

    # Bloqueo backend: si tiene factura, no se puede borrar
    if Factura.objects.filter(numero_factura=cuenta.numero_factura).exists():
        messages.error(request, '❌ No se puede eliminar: tiene una factura asociada. Elimina primero la factura.')
        return redirect('cuentas_por_pagar')

    if request.method == 'POST':
        concepto = cuenta.concepto
        cuenta.delete()
        messages.success(request, f'✅ Cuenta "{concepto}" eliminada correctamente')
        return redirect('cuentas_por_pagar')

    return redirect('confirmar_eliminar_cuenta_por_pagar', cuenta_id=cuenta_id)


@login_required
def confirmar_eliminar_cuenta_por_cobrar(request, cuenta_id):
    """Muestra página de confirmación para eliminar una cuenta por cobrar"""
    cuenta = get_object_or_404(CuentaPorCobrar, id=cuenta_id)

    tiene_factura = Factura.objects.filter(numero_factura=cuenta.numero_factura).exists()

    context = {
        'cuenta': cuenta,
        'tiene_factura': tiene_factura,
    }
    return render(request, 'finanzas/confirmar_eliminar_cuenta_cobrar.html', context)


@login_required
def eliminar_cuenta_por_cobrar(request, cuenta_id):
    """Elimina definitivamente una cuenta por cobrar (solo si no tiene factura)"""
    cuenta = get_object_or_404(CuentaPorCobrar, id=cuenta_id)

    if Factura.objects.filter(numero_factura=cuenta.numero_factura).exists():
        messages.error(request, '❌ No se puede eliminar: tiene una factura asociada. Elimina primero la factura.')
        return redirect('cuentas_por_cobrar')

    if request.method == 'POST':
        concepto = cuenta.concepto
        cuenta.delete()
        messages.success(request, f'✅ Cuenta "{concepto}" eliminada correctamente')
        return redirect('cuentas_por_cobrar')

    return redirect('confirmar_eliminar_cuenta_por_cobrar', cuenta_id=cuenta_id)

# ============================================
# EXPORTAR A EXCEL
# ============================================

from core.utils import exportar_excel
from datetime import date


@login_required
def exportar_facturas_excel(request):
    """Exporta facturas a Excel."""
    facturas = Factura.objects.exclude(estado_pago='ANULADO').order_by('-fecha')

    # Filtros opcionales
    desde = request.GET.get('desde')
    hasta = request.GET.get('hasta')
    tipo = request.GET.get('tipo')

    if desde:
        facturas = facturas.filter(fecha__gte=desde)
    if hasta:
        facturas = facturas.filter(fecha__lte=hasta)
    if tipo:
        facturas = facturas.filter(tipo=tipo)

    cabeceras = [
        'Nº Factura', 'Tipo', 'Fecha', 'Vencimiento',
        'Cliente/Proveedor', 'Categoría', 'Base (€)', 'IVA (%)', 'IVA (€)', 'Total (€)',
        'Estado Pago', 'Notas'
    ]

    filas = []
    for f in facturas:
        entidad = f.cliente.nombre if f.cliente else (f.proveedor.nombre if f.proveedor else '—')
        filas.append([
            f.numero_factura,
            f.get_tipo_display(),
            f.fecha.strftime('%d/%m/%Y') if f.fecha else '',
            f.fecha_vencimiento.strftime('%d/%m/%Y') if f.fecha_vencimiento else '',
            entidad,
            f.get_categoria_display(),
            float(f.base),
            float(f.iva_porcentaje),
            float(f.iva),
            float(f.total),
            f.get_estado_pago_display(),
            f.notas or '',
        ])

    hoy = date.today().strftime('%Y%m%d')
    nombre = f'Facturas_{hoy}.xlsx'

    return exportar_excel(nombre, cabeceras, filas, titulo_hoja="Facturas")

# ============================================
# LIBRO DE GASTOS E INGRESOS
# ============================================

@login_required
def libro_gastos_ingresos(request):
    """Libro mayor con ingresos y gastos separados, con filtros."""
    from django.db.models import Sum
    from decimal import Decimal

    facturas = Factura.objects.exclude(estado_pago='ANULADO').order_by('-fecha')

    # Filtros
    categoria = request.GET.get('categoria', '')
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')

    if categoria:
        facturas = facturas.filter(categoria=categoria)
    if desde:
        facturas = facturas.filter(fecha__gte=desde)
    if hasta:
        facturas = facturas.filter(fecha__lte=hasta)

    # Separar en dos listas
    facturas_ingresos = facturas.filter(tipo='INGRESO')
    facturas_gastos = facturas.filter(tipo='GASTO')

    # Totales
    total_ingresos = facturas_ingresos.aggregate(t=Sum('total'))['t'] or Decimal('0')
    total_gastos = facturas_gastos.aggregate(t=Sum('total'))['t'] or Decimal('0')
    beneficio = total_ingresos - total_gastos
    num_facturas = facturas.count()

    context = {
        'facturas_ingresos': facturas_ingresos,
        'facturas_gastos': facturas_gastos,
        'total_ingresos': total_ingresos,
        'total_gastos': total_gastos,
        'beneficio': beneficio,
        'num_facturas': num_facturas,
        'categorias': Factura.CATEGORIAS,
        'categoria_filtro': categoria,
        'desde_filtro': desde,
        'hasta_filtro': hasta,
    }
    return render(request, 'finanzas/libro_gastos_ingresos.html', context)

@login_required
def exportar_libro_excel(request):
    """Exporta el libro con hojas por trimestre + hoja Resumen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from datetime import date

    # Filtros
    categoria = request.GET.get('categoria', '')
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    year = request.GET.get('year', date.today().year)

    facturas = Factura.objects.exclude(estado_pago='ANULADO').order_by('fecha')
    if categoria:
        facturas = facturas.filter(categoria=categoria)
    if desde:
        facturas = facturas.filter(fecha__gte=desde)
    if hasta:
        facturas = facturas.filter(fecha__lte=hasta)

    wb = Workbook()
    wb.remove(wb.active)

    cabeceras = ['Fecha', 'Nº Factura', 'Entidad', 'Categoría', 'Base (€)', 'IVA (%)', 'IVA (€)', 'Total (€)', 'Estado']

    estilo_cabecera = Font(bold=True, color='FFFFFF')
    fondo_ingresos = PatternFill('solid', fgColor='2D6A4F')
    fondo_gastos = PatternFill('solid', fgColor='DC3545')

    # --- Hoja RESUMEN ---
    ws_resumen = wb.create_sheet(title="Resumen")
    ws_resumen['A1'] = f'RESUMEN ANUAL {year}'
    ws_resumen['A1'].font = Font(bold=True, size=14)
    ws_resumen.append([])
    ws_resumen.append(['Tipo', 'Total (€)'])
    for cell in ws_resumen[3]:
        cell.font = estilo_cabecera
        cell.fill = fondo_ingresos

    total_ing_año = Decimal('0')
    total_gas_año = Decimal('0')

    # --- Hojas por trimestre ---
    for trimestre in range(1, 5):
        mes_inicio = (trimestre - 1) * 3 + 1
        mes_fin = mes_inicio + 2

        facturas_trim = facturas.filter(fecha__month__gte=mes_inicio, fecha__month__lte=mes_fin)
        facturas_ing = facturas_trim.filter(tipo='INGRESO')
        facturas_gas = facturas_trim.filter(tipo='GASTO')

        total_ing = facturas_ing.aggregate(t=Sum('total'))['t'] or Decimal('0')
        total_gas = facturas_gas.aggregate(t=Sum('total'))['t'] or Decimal('0')
        total_ing_año += total_ing
        total_gas_año += total_gas

        ws = wb.create_sheet(title=f"T{trimestre} {year}")

        # Sección INGRESOS
        ws.append([f'INGRESOS T{trimestre}'])
        ws['A1'].font = Font(bold=True, size=12)
        ws.append(cabeceras)
        for cell in ws[2]:
            cell.font = estilo_cabecera
            cell.fill = fondo_ingresos
        for f in facturas_ing:
            entidad = f.entidad_nombre
            ws.append([f.fecha.strftime('%d/%m/%Y'), f.numero_factura, entidad, f.get_categoria_display(),
                       float(f.base), float(f.iva_porcentaje), float(f.iva), float(f.total), f.get_estado_pago_display()])
        ws.append(['TOTAL INGRESOS', '', '', '', '', '', '', float(total_ing), ''])
        ws.append([])

        # Sección GASTOS
        fila_gastos_titulo = ws.max_row + 1
        ws.append([f'GASTOS T{trimestre}'])
        ws.cell(row=fila_gastos_titulo, column=1).font = Font(bold=True, size=12)
        fila_gastos_cab = ws.max_row + 1
        ws.append(cabeceras)
        for cell in ws[fila_gastos_cab]:
            cell.font = estilo_cabecera
            cell.fill = fondo_gastos
        for f in facturas_gas:
            entidad = f.entidad_nombre
            ws.append([f.fecha.strftime('%d/%m/%Y'), f.numero_factura, entidad, f.get_categoria_display(),
                       float(f.base), float(f.iva_porcentaje), float(f.iva), float(f.total), f.get_estado_pago_display()])
        ws.append(['TOTAL GASTOS', '', '', '', '', '', '', float(total_gas), ''])
        ws.append([])
        ws.append([f'BENEFICIO T{trimestre}', '', '', '', '', '', '', float(total_ing - total_gas), ''])

    # Rellenar resumen
    ws_resumen.append(['Ingresos', float(total_ing_año)])
    ws_resumen.append(['Gastos', float(total_gas_año)])
    ws_resumen.append(['Beneficio', float(total_ing_año - total_gas_año)])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Libro_{year}.xlsx"'
    wb.save(response)
    return response

@login_required
def anular_factura(request, factura_id):
    """Anula una factura (no la elimina, la marca como ANULADA)"""
    factura = get_object_or_404(Factura, id=factura_id)

    if request.method == 'POST':
        numero = factura.numero_factura
        factura.estado_pago = 'ANULADO'
        factura.save()
        
        # Anular también sus cuentas asociadas
        # 1. Por numero_factura exacto
        CuentaPorCobrar.objects.filter(numero_factura=numero).update(estado='ANULADO')
        CuentaPorPagar.objects.filter(numero_factura=numero).update(estado='ANULADO')
        
        # 2. Por concepto que contenga el número
        CuentaPorCobrar.objects.filter(concepto__icontains=numero).update(estado='ANULADO')
        CuentaPorPagar.objects.filter(concepto__icontains=numero).update(estado='ANULADO')
        
        # 3. Por pedido asociado (si la factura tiene pedido)
        if factura.pedido:
            CuentaPorCobrar.objects.filter(concepto__icontains=factura.pedido.numero_pedido).update(estado='ANULADO')
            CuentaPorPagar.objects.filter(concepto__icontains=factura.pedido.numero_pedido).update(estado='ANULADO')
        
        messages.success(request, f'✅ Factura {numero} y sus cuentas asociadas anuladas')
        return redirect('lista_facturas')

    return redirect('lista_facturas')

@login_required
def lista_comisiones(request):
    """Lista todas las comisiones agrupadas por comercial"""
    comerciales = Comercial.objects.filter(activo=True).prefetch_related('comision_set')
    
    datos_comerciales = []
    for com in comerciales:
        comisiones = com.comision_set.all().select_related('pedido')
        total_comision = sum(c.importe_comision for c in comisiones)
        
        datos_comerciales.append({
            'comercial': com,
            'comisiones': comisiones,
            'total_comision': total_comision,
            'num_pedidos': len(comisiones),
        })
        
    # Ordenar por el que más ha ganado
    datos_comerciales = sorted(datos_comerciales, key=lambda x: x['total_comision'], reverse=True)

    context = {
        'datos_comerciales': datos_comerciales,
    }
    return render(request, 'finanzas/lista_comisiones.html', context)